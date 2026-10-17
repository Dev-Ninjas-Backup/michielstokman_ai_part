import io
import json
import uuid
import pytest
from unittest.mock import MagicMock, patch

from app.utils.document_parser import extract_text_from_file
from app.services.service_bulk_import import (
    _normalize_story_item,
    _clean_json_response,
    parse_and_split_document_text,
    ingest_bulk_stories,
    CANONICAL_GROWTH_AREAS,
    CANONICAL_LIFE_PHASES,
)
from app.model.story import StoryType, AssetReviewStatus, ModerationStatus


def test_document_parser_txt():
    content = "Story 1\nThis is a confession.\n\nStory 2\nThis is a meditation."
    raw = content.encode("utf-8")
    extracted = extract_text_from_file(raw, "document.txt")
    assert "Story 1" in extracted
    assert "Story 2" in extracted


def test_document_parser_docx():
    import docx
    doc = docx.Document()
    doc.add_paragraph("First Story Headline")
    doc.add_paragraph("Narrative content of the story.")
    buf = io.BytesIO()
    doc.save(buf)

    extracted = extract_text_from_file(buf.getvalue(), "stories.docx")
    assert "First Story Headline" in extracted
    assert "Narrative content" in extracted


def test_normalize_story_item_canonical_mappings():
    raw_item = {
        "title": "A Walk by the Sea",
        "story_type": "meditation",
        "story_text": "I walked slowly along the morning shore.",
        "first_name": "Sarah",
        "gender": "Female",
        "age": "29",
        "city": "Brighton",
        "country": "UK",
        "growth_areas": ["fear and freedom", "letting go", "unrecognized area"],
        "life_phase": "deepening",
        "tags": ["Ocean", "Calm", "MORNING"],
        "high_intensity": False,
        "hero_hook": "The salt in the air reminded me of another time.",
        "hero_tagline": "Finding peace in the morning tide.",
    }

    norm = _normalize_story_item(raw_item)
    assert norm["title"] == "A Walk by the Sea"
    assert norm["story_type"] == "meditation"
    assert norm["age"] == 29
    assert norm["gender"] == "female"
    assert "Fear & Freedom" in norm["growth_areas"]
    assert "Letting Go" in norm["growth_areas"]
    assert norm["life_phase"] == "Deepening"
    assert norm["tags"] == ["ocean", "calm", "morning"]


def test_clean_json_response():
    markdown_wrapped = "```json\n{\"stories\": [{\"title\": \"One\"}]}\n```"
    assert _clean_json_response(markdown_wrapped) == "{\"stories\": [{\"title\": \"One\"}]}"


@patch("app.services.service_bulk_import.get_story_llm")
def test_parse_and_split_document_text(mock_get_llm):
    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(content=json.dumps({
        "stories": [
            {
                "title": "Story Alpha",
                "story_type": "confession",
                "story_text": "First story narrative text here.",
                "first_name": "Elena",
                "growth_areas": ["Self-Acceptance"],
                "life_phase": "Discovering",
            },
            {
                "title": "Story Beta",
                "story_type": "meditation",
                "story_text": "Second story narrative text here.",
                "first_name": "Lucas",
                "growth_areas": ["Presence"],
                "life_phase": "Building",
            }
        ]
    }))
    mock_get_llm.return_value = mock_llm

    document_text = "Story Alpha: First story...\n\nStory Beta: Second story..."
    results = parse_and_split_document_text(document_text)

    assert len(results) == 2
    assert results[0]["title"] == "Story Alpha"
    assert results[0]["story_type"] == "confession"
    assert results[1]["title"] == "Story Beta"
    assert results[1]["story_type"] == "meditation"


def test_ingest_bulk_stories_creates_rows_and_queues_tasks():
    from types import SimpleNamespace

    stories_data = [
        {
            "title": "Ingested Story 1",
            "story_type": "confession",
            "story_text": "Story one text narrative.",
            "first_name": "Marcus",
            "gender": "male",
            "growth_areas": ["Fear & Freedom"],
            "life_phase": "Discovering",
            "hero_hook": "Hook 1",
            "hero_tagline": "Tagline 1",
        }
    ]

    db = MagicMock()
    mock_story = SimpleNamespace(
        id=uuid.uuid4(),
        story_number=101,
        title="Ingested Story 1",
        story_type=StoryType.confession,
        cover_status=AssetReviewStatus.in_progress,
        voice_status=AssetReviewStatus.missing,
        content_status=AssetReviewStatus.ready_for_review,
        moderation_status=ModerationStatus.pending,
    )

    with patch("app.data.story.create_story", return_value=mock_story):
        bg_tasks = MagicMock()
        res = ingest_bulk_stories(
            db=db,
            stories_data=stories_data,
            admin_id=str(uuid.uuid4()),
            background_tasks=bg_tasks,
            generate_covers=True,
            generate_audio=False,
        )

        assert len(res) == 1
        assert res[0]["title"] == "Ingested Story 1"
        assert res[0]["cover_status"] == "in_progress"
        assert res[0]["voice_status"] == "missing"
        # Background task for cover should have been scheduled
        assert bg_tasks.add_task.called
